#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import fitz
from openpyxl import load_workbook


DEFAULT_SAMPLE_ROOT = Path("/Users/xiexiaoxiong/Documents/Patent项目配套/实例专利")
DEFAULT_ORACLE = Path(__file__).resolve().parents[1] / "fixtures" / "sample-oracle.json"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def inspect_file(path: Path, root: Path, oracle: dict[str, Any]) -> dict[str, Any]:
    suffix = path.suffix.lower()
    item: dict[str, Any] = {
        "relative_path": str(path.relative_to(root)),
        "file_name": path.name,
        "suffix": suffix,
        "byte_size": path.stat().st_size,
        "sha256": file_sha256(path),
        "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
    }
    if suffix == ".pdf":
        item["category"] = "patent_pdf"
        document = fitz.open(path)
        try:
            item["page_count"] = document.page_count
            item["text_layer_pages"] = sum(
                bool(document.load_page(index).get_text().strip())
                for index in range(document.page_count)
            )
        finally:
            document.close()
        expected = oracle.get("patents", {}).get(path.name)
        item["oracle"] = expected
        item["oracle_status"] = "matched" if expected else "missing"
    elif suffix == ".xlsx":
        item["category"] = "auxiliary_product_xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            item["sheet_names"] = workbook.sheetnames
            item["sheet_count"] = len(workbook.sheetnames)
        finally:
            workbook.close()
    else:
        item["category"] = "unclassified"
    return item


def build_manifest(root: Path, oracle_path: Path) -> dict[str, Any]:
    if not root.is_dir():
        raise FileNotFoundError(root)
    oracle = json.loads(oracle_path.read_text(encoding="utf-8"))
    files = [
        inspect_file(path, root, oracle)
        for path in sorted(root.rglob("*"), key=lambda item: str(item).lower())
        if path.is_file()
    ]
    patent_names = {item["file_name"] for item in files if item["category"] == "patent_pdf"}
    expected_names = set(oracle.get("patents", {}))
    return {
        "schema_version": "1.0",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "sample_root": str(root.resolve()),
        "file_count": len(files),
        "patent_pdf_count": sum(item["category"] == "patent_pdf" for item in files),
        "auxiliary_xlsx_count": sum(
            item["category"] == "auxiliary_product_xlsx" for item in files
        ),
        "missing_oracle_files": sorted(expected_names - patent_names),
        "unexpected_patent_files": sorted(patent_names - expected_names),
        "files": files,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=DEFAULT_SAMPLE_ROOT)
    parser.add_argument("--oracle", type=Path, default=DEFAULT_ORACLE)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    manifest = build_manifest(args.root, args.oracle)
    encoded = json.dumps(manifest, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded + "\n", encoding="utf-8")
    else:
        print(encoded)
    if args.check:
        if manifest["missing_oracle_files"]:
            return 2
        if manifest["unexpected_patent_files"]:
            return 4
        if manifest["patent_pdf_count"] < 1:
            return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
